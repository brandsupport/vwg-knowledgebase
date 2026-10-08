#!/usr/bin/env python3
"""
Reads data/articles.xlsx (Microsoft Forms responses) and converts approved
rows into knowledgebase articles.

The sheet is found automatically: any worksheet whose header row contains a
"Title" column and a "Body" column is used, whatever the sheet is called.

Columns (header names are case-insensitive):
    ID            added by Forms; keeps article ids stable between builds
    Title         required
    Category      one of the names on the Categories sheet
    Tags          optional, comma/semicolon separated
    Body          required; see FORMATTING below
    Images        optional, comma/semicolon separated filenames (e.g.
                  "screenshot1.png, screenshot2.png"). Each file must be
                  uploaded to data/images/ in the repo - this column just
                  says which of those files belong to this article.
    Source link   optional, http(s) only
    Status        only rows marked Approved are published.
                  If the sheet has no Status column at all, every row is
                  published (a warning is printed).

FORMATTING inside Body:
    blank line          new paragraph
    - text              bullet
    1. text             numbered step
    ## Heading          section heading
    Note: text          highlighted warning box
    **bold**            bold
    emails and bulletin codes (e.g. VW0948-26) are styled automatically.
All text is HTML-escaped, so nothing typed into the form can inject markup.
"""
import os
from openpyxl import load_workbook
from article_formatting import (
    ALIASES, APPROVED, ID_OFFSET, category_lookup, load_local_image, build_article,
)


def _find_sheet(wb):
    for ws in wb.worksheets:
        for r in range(1, 11):
            headers = [str(c.value).strip().lower() if c.value is not None else '' for c in ws[r]]
            cols = {}
            for key, names in ALIASES.items():
                for i, h in enumerate(headers):
                    if h in names:
                        cols[key] = i
                        break
            if 'title' in cols and 'body' in cols:
                return ws, r, cols
    return None, None, None


def load_excel_articles(path, categories, images_dir=None):
    """Returns (articles, warnings)."""
    warnings = []
    wb = load_workbook(path, data_only=True)
    ws, header_row, cols = _find_sheet(wb)
    if ws is None:
        return [], [f'{path}: no sheet with "Title" and "Body" columns was found']

    if 'status' not in cols:
        warnings.append('No Status column found - every row will be published. '
                        'Add a Status column to control what goes live.')

    if images_dir is None:
        images_dir = os.path.join(os.path.dirname(os.path.abspath(path)), 'images')

    cat_lut = category_lookup(categories)
    cat_short = {c['id']: c['short'] for c in categories}
    fallback_cat = cat_lut.get('reference')

    def cell(row, key):
        i = cols.get(key)
        if i is None or i >= len(row):
            return ''
        v = row[i].value
        return '' if v is None else str(v).strip()

    articles = []
    for offset, row in enumerate(ws.iter_rows(min_row=header_row + 1), start=1):
        title, body = cell(row, 'title'), cell(row, 'body')
        if not title and not body:
            continue
        if 'status' in cols and cell(row, 'status').lower() not in APPROVED:
            continue
        if not title or not body:
            warnings.append(f'Row {header_row + offset}: skipped (Title and Body are both required)')
            continue

        row_label = f'Row {header_row + offset}'
        raw_cat = cell(row, 'category')
        cat = cat_lut.get(raw_cat.lower())
        if cat is None:
            cat = fallback_cat
            warnings.append(f'{row_label}: category "{raw_cat}" not recognised - filed under Reference')

        raw_id = cell(row, 'id')
        try:
            art_id = ID_OFFSET + int(float(raw_id))
        except ValueError:
            art_id = ID_OFFSET + header_row + offset

        images = []
        for name in cell(row, 'images').replace(';', ',').split(','):
            name = name.strip()
            if not name:
                continue
            uri = load_local_image(name, images_dir, warnings, row_label)
            if uri:
                images.append(uri)

        articles.append(build_article(
            art_id=art_id, cat=cat, cat_short_map=cat_short,
            title=title, body=body, user_tags_raw=cell(row, 'tags'),
            images=images, source_url=cell(row, 'source'), source='excel',
        ))

    articles.sort(key=lambda a: a['id'], reverse=True)  # newest first
    seen = set()
    for a in articles:
        if a['id'] in seen:
            warnings.append(f'Duplicate ID {a["id"] - ID_OFFSET} in the sheet - check the ID column')
        seen.add(a['id'])
    return articles, warnings


if __name__ == '__main__':
    import json, sys
    root = os.path.dirname(os.path.abspath(__file__))
    data = json.load(open(os.path.join(root, 'data', 'site_data.json'), encoding='utf-8'))
    arts, warns = load_excel_articles(sys.argv[1] if len(sys.argv) > 1 else os.path.join(root, 'data', 'articles.xlsx'), data['categories'])
    print(f'{len(arts)} article(s) would be published')
    for a in arts:
        print(f'  #{a["id"]}  [{a["tags"][0]}]  {a["title"]}')
    for w in warns:
        print('  WARNING:', w)
